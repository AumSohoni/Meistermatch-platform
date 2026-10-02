import sys
from docx import Document
sys.stdout.reconfigure(encoding="utf-8")

doc = Document(r"C:\Users\aum11\Downloads\6_Sem_Study\Bachelor_Thesis\thesis_5th_semester_report_v11.docx")
full = "\n".join(p.text for p in doc.paragraphs)
low = full.lower()

checks = {
    "supplementary": "supplementary" in low,
    "markdown hashes": "##" in full,
    "TODO/TBD": ("todo" in low) or ("tbd" in low),
    "framework?s": "framework?s" in full,
    "system?s": "system?s" in full,
    "Critically,": "critically," in low,
    "comprehensive": "comprehensive" in low,
    "Lorem ipsum": "lorem" in low,
    "draft appendix": "draft" in low,
}
for k, v in checks.items():
    print(f"{k}: {'HIT' if v else 'ok'}")

xml = doc.settings.element.xml if hasattr(doc.settings.element, "xml") else ""
print("updateFields present:", "updateFields" in xml)
print("total paragraphs:", len(doc.paragraphs))
print("tables:", len(doc.tables))

for sname in ["Heading 1", "Heading 2", "Heading 3"]:
    st = doc.styles[sname]
    print(
        sname,
        "align=", st.paragraph_format.alignment,
        "size=", st.font.size.pt,
        "bold=", st.font.bold,
        "font=", st.font.name,
    )

p = doc.paragraphs[97]
print("ref[97] left_indent:", p.paragraph_format.left_indent,
      "first_line_indent:", p.paragraph_format.first_line_indent,
      "align:", p.alignment)

print("H1:", sum(1 for p in doc.paragraphs if p.style.name == "Heading 1"))
print("H2:", sum(1 for p in doc.paragraphs if p.style.name == "Heading 2"))
print("H3:", sum(1 for p in doc.paragraphs if p.style.name == "Heading 3"))

refs = [p for p in doc.paragraphs if p.paragraph_format.left_indent is not None and p.paragraph_format.first_line_indent is not None and p.text.strip()]
print("reference entries:", len(refs))